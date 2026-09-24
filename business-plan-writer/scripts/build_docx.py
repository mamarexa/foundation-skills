#!/usr/bin/env python3
"""Render a business-plan spec (JSON) into a styled Word document.

    python build_docx.py plan.json business-plan.docx [--facts facts.json] [--pdf]

Needs python-docx and matplotlib (both preinstalled in Claude's code-execution sandbox;
elsewhere: pip install python-docx matplotlib). --pdf additionally converts with LibreOffice
(soffice) when it's installed, for visual QA.

The spec format is documented in ../references/spec-format.md. Every {{path|format}} token is
filled from the facts file (the Foundation engine's own output). See foundation_facts.py.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from foundation_facts import Facts, FactsError  # noqa: E402

DEFAULT_THEME = {
    "heading_font": "Georgia",
    "body_font": "Calibri",
    "ink": "111111",
    "muted": "787774",
    "border": "D9D8D4",
    "fill": "F7F6F3",
    "accent": "1F6C9F",
}

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
NUMERIC = re.compile(r"^[−\-+]?[$€£¥₹₺A-Z]*\$?[\d,.]+[%×kMB]?( ?[A-Z]{3})?( years| months)?$")


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


class Builder:
    def __init__(self, spec: dict, facts: Facts):
        self.spec = spec
        self.facts = facts
        self.theme = {**DEFAULT_THEME, **spec.get("theme", {})}
        self.doc = Document()
        self.figure_no = 0
        self._setup_styles()

    # --- setup ------------------------------------------------------------------------------

    def _setup_styles(self) -> None:
        t = self.theme
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = (Cm(21.0), Cm(29.7)) if self.spec.get("paper", "A4") == "A4" else (Cm(21.59), Cm(27.94))
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.3))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

        normal = self.doc.styles["Normal"]
        normal.font.name = t["body_font"]
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), t["body_font"])
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = rgb(t["ink"])
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.2

        for name, size, before, after in (("Title", 30, 0, 6), ("Heading 1", 20, 0, 10), ("Heading 2", 14, 14, 4), ("Heading 3", 11.5, 10, 2)):
            st = self.doc.styles[name]
            st.font.name = t["heading_font"]
            st.element.rPr.rFonts.set(qn("w:eastAsia"), t["heading_font"])
            st.font.size = Pt(size)
            st.font.bold = name == "Heading 3"
            st.font.color.rgb = rgb(t["ink"])
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
            p_pr = st.element.find(qn("w:pPr"))
            if p_pr is not None and p_pr.find(qn("w:pBdr")) is not None:
                p_pr.remove(p_pr.find(qn("w:pBdr")))  # Word's default Title has a coloured rule under it

    # --- text helpers -----------------------------------------------------------------------

    def text(self, value: str) -> str:
        return self.facts.fill(str(value))

    def add_runs(self, paragraph, value: str, size: float | None = None, color: str | None = None) -> None:
        for part in INLINE.split(self.text(value)):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*"):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(part)
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = rgb(color)

    # --- document parts ---------------------------------------------------------------------

    def cover(self) -> None:
        s, t = self.spec, self.theme
        for _ in range(6):
            self.doc.add_paragraph()
        if s.get("kicker"):
            p = self.doc.add_paragraph()
            self.add_runs(p, s["kicker"].upper(), size=9, color=t["muted"])
        self.doc.add_paragraph(self.text(s["title"]), style="Title")
        if s.get("subtitle"):
            p = self.doc.add_paragraph()
            self.add_runs(p, s["subtitle"], size=13, color=t["muted"])
        for _ in range(10):
            self.doc.add_paragraph()
        for line in (s.get("prepared_by"), s.get("date"), s.get("confidentiality")):
            if line:
                p = self.doc.add_paragraph()
                self.add_runs(p, line, size=9.5, color=t["muted"])
                p.paragraph_format.space_after = Pt(2)

    def contents(self) -> None:
        self.page_break()
        self.doc.add_paragraph("Contents", style="Heading 1")
        for i, section in enumerate(self.spec["sections"], 1):
            p = self.doc.add_paragraph()
            self.add_runs(p, f"{i}.  {section['heading']}")
            p.paragraph_format.space_after = Pt(3)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def header_footer(self) -> None:
        # Body pages get the plan's title in the header and a page number in the footer; the
        # cover (first page of the first section) stays clean.
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        self.add_runs(hp, self.spec.get("short_title") or self.spec["title"], size=8.5, color=self.theme["muted"])
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = fp.add_run()
        run.font.size = Pt(8.5)
        run.font.color.rgb = rgb(self.theme["muted"])
        for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    # --- blocks -----------------------------------------------------------------------------

    def block(self, b: dict) -> None:
        kind = b.get("type", "paragraph")
        handler = getattr(self, f"b_{kind}", None)
        if handler is None:
            raise ValueError(f'unknown block type "{kind}"')
        handler(b)

    def b_paragraph(self, b: dict) -> None:
        self.add_runs(self.doc.add_paragraph(), b["text"])

    def b_heading(self, b: dict) -> None:
        self.doc.add_paragraph(self.text(b["text"]), style="Heading 3" if b.get("level", 2) >= 3 else "Heading 2")

    def b_bullets(self, b: dict) -> None:
        for item in b["items"]:
            p = self.doc.add_paragraph(style="List Number" if b.get("numbered") else "List Bullet")
            self.add_runs(p, item)
            p.paragraph_format.space_after = Pt(3)

    def b_callout(self, b: dict) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        shade(cell, self.theme["fill"])
        set_cell_borders(cell, left=self.theme["ink"], width=18)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        if b.get("title"):
            self.add_runs(cell.paragraphs[0], f"**{b['title']}**")
            p = cell.add_paragraph()
        else:
            p = cell.paragraphs[0]
        self.add_runs(p, b["text"])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def b_kpis(self, b: dict) -> None:
        items = b["items"]
        table = self.doc.add_table(rows=2, cols=len(items))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, item in enumerate(items):
            top, bottom = table.cell(0, i), table.cell(1, i)
            self.add_runs(top.paragraphs[0], item["value"], size=17)
            top.paragraphs[0].runs[0].font.name = self.theme["heading_font"]
            self.add_runs(bottom.paragraphs[0], item["label"], size=8.5, color=self.theme["muted"])
            for c in (top, bottom):
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_cell_borders(bottom, bottom=self.theme["border"])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def b_table(self, b: dict) -> None:
        self.table(b["columns"], b["rows"], caption=b.get("caption"), total_row=b.get("total_row", False))

    def b_projection_table(self, b: dict) -> None:
        header, rows = self.facts.projection_rows()
        self.table(header, rows, caption=b.get("caption", "Source: Foundation feasibility engine."), filled=True)

    def b_breakdown_table(self, b: dict) -> None:
        header, rows = self.facts.breakdown_rows(b["of"])
        self.table(header, rows, caption=b.get("caption", "Source: Foundation feasibility engine."), total_row=True, filled=True)

    def b_scenarios_table(self, b: dict) -> None:
        header, rows = self.facts.scenario_rows()
        self.table(header, rows, caption=b.get("caption", "Source: Foundation scenario planning (Worst / Base / Best)."), filled=True)

    def b_inputs_table(self, b: dict) -> None:
        header, rows = self.facts.input_rows(b["category"])
        if rows:
            self.table(header, rows, caption=b.get("caption"), filled=True)

    def table(self, columns, rows, caption=None, total_row=False, filled=False) -> None:
        t = self.theme
        header = [self.text(c) for c in columns]
        body = [[cell if filled else self.text(cell) for cell in row] for row in rows]
        table = self.doc.add_table(rows=1 + len(body), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        numeric_cols = [all(NUMERIC.match(str(r[i]).strip() or "0") for r in body) for i in range(len(header))]
        for r, values in enumerate([header, *body]):
            is_total = total_row and r == len(body)
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = r < len(body)  # keep a table on one page where it fits
                if numeric_cols[c] and c > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if r == 0:
                    self.add_runs(p, str(value), size=9, color=t["muted"])
                    shade(cell, t["fill"])
                    set_cell_borders(cell, bottom=t["border"])
                else:
                    self.add_runs(p, f"**{value}**" if is_total else str(value), size=9.5)
                    set_cell_borders(cell, bottom=t["ink"] if is_total else t["border"], top=t["ink"] if is_total else None)
        if caption:
            p = self.doc.add_paragraph()
            self.add_runs(p, caption, size=8.5, color=t["muted"])
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def b_chart(self, b: dict) -> None:
        png = render_chart(b, self.facts, self.theme)
        self.doc.add_picture(png, width=Cm(b.get("width_cm", 15.5)))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.figure_no += 1
        p = self.doc.add_paragraph()
        self.add_runs(p, f"Figure {self.figure_no}. {b.get('title', '')}".strip(), size=8.5, color=self.theme["muted"])

    def b_pagebreak(self, _b: dict) -> None:
        self.page_break()

    # --- whole document ---------------------------------------------------------------------

    def build(self) -> Document:
        self.cover()
        if self.spec.get("contents", True):
            self.contents()
        for i, section in enumerate(self.spec["sections"], 1):
            self.page_break()
            self.doc.add_paragraph(f"{i}. {self.text(section['heading'])}", style="Heading 1")
            for b in section.get("blocks", []):
                self.block(b)
        self.header_footer()
        return self.doc


# --- charts ---------------------------------------------------------------------------------

def render_chart(b: dict, facts: Facts, theme: dict) -> io.BytesIO:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter

    if b.get("from"):  # engine series, e.g. {"from": "results.projection", "x": "year", "series": [{"field": "revenue", "name": "Revenue"}]}
        labels = [f"Y{int(v)}" for v in facts.series(b["from"], b.get("x", "year"))]
        series = [{"name": s["name"], "values": facts.series(b["from"], s["field"]), "kind": s.get("kind")} for s in b["series"]]
    else:
        labels = [str(x) for x in b["labels"]]
        series = b["series"]

    palette = [f"#{theme['ink']}", f"#{theme['accent']}", "#A3A29C", "#C9C7C0"]
    fig, ax = plt.subplots(figsize=(b.get("width_in", 7.2), b.get("height_in", 3.2)), dpi=200)
    kind = b.get("kind", "bar")
    n = len(series)
    width = 0.8 / max(1, sum(1 for s in series if (s.get("kind") or kind) == "bar"))
    bar_i = 0
    xs = range(len(labels))
    for i, s in enumerate(series):
        color = palette[i % len(palette)]
        if (s.get("kind") or kind) == "line":
            ax.plot(list(xs), s["values"], color=color, linewidth=2, marker="o", markersize=3.5, label=s["name"])
        else:
            offset = (bar_i - (n - 1) / 2) * width if n > 1 else 0
            ax.bar([x + offset for x in xs], s["values"], width=width * 0.92, color=color, label=s["name"])
            bar_i += 1
    ax.set_xticks(list(xs), labels)
    money = b.get("unit", "money") == "money"
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: facts.money(v, compact=True) if money else f"{v:,.0f}"))
    ax.axhline(0, color="#999999", linewidth=0.8)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("#BBBBBB")
    ax.tick_params(colors="#555555", labelsize=8, length=0)
    ax.grid(axis="y", color="#EEEEEE", linewidth=0.8)
    ax.set_axisbelow(True)
    if n > 1:
        ax.legend(frameon=False, fontsize=8, loc="upper left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


# --- low-level table styling ----------------------------------------------------------------

def shade(cell, hex_: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tc_pr.append(shd)


def set_cell_borders(cell, width: int = 6, **sides) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        color = sides.get(side)
        el.set(qn("w:val"), "single" if color else "nil")
        if color:
            el.set(qn("w:sz"), str(width))
            el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


# --- entry point ----------------------------------------------------------------------------

def verify(path: Path) -> list[str]:
    """Re-open the file and look for anything that shouldn't ship."""
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    problems = []
    if "{{" in text or "}}" in text:
        problems.append("unresolved {{token}} left in the document")
    for marker in ("TODO", "TBD", "lorem ipsum", "[insert"):
        if marker.lower() in text.lower():
            problems.append(f'placeholder text "{marker}" found')
    return problems


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("out")
    ap.add_argument("--facts", help="JSON with the Foundation engine's output (see foundation_facts.py)")
    ap.add_argument("--pdf", action="store_true", help="also export a PDF with LibreOffice, for visual QA")
    args = ap.parse_args()

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    try:
        doc = Builder(spec, Facts.load(args.facts)).build()
    except FactsError as e:
        print(f"ERROR (numbers): {e}", file=sys.stderr)
        return 2
    out = Path(args.out)
    doc.save(str(out))
    problems = verify(out)
    for p in problems:
        print(f"WARNING: {p}", file=sys.stderr)
    print(f"wrote {out}")

    if args.pdf:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            print("note: LibreOffice not installed, skipped PDF export", file=sys.stderr)
        else:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out.parent), str(out)],
                           check=True, capture_output=True, timeout=180)
            print(f"wrote {out.with_suffix('.pdf')}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
