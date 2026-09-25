"""Smoke tests for the document skills' renderers.

    python skills/tests/test_scripts.py

Builds the café fixtures (whose facts come from the real Foundation engine, see
scripts/make-skill-fixture-facts.ts in the app repo) and checks the guarantees the skills rely on.
"""

import filecmp
import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures"
DOCX = ROOT / "business-plan-writer" / "scripts" / "build_docx.py"
PPTX = ROOT / "pitch-deck-writer" / "scripts" / "build_pptx.py"
failures = []


def check(name, cond, detail=""):
    print(("ok   " if cond else "FAIL ") + name + (f": {detail}" if detail and not cond else ""))
    if not cond:
        failures.append(name)


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


for script in ("foundation_facts.py", "writing_lint.py", "check_numbers.py"):
    check(
        f"{script} copies are identical",
        filecmp.cmp(ROOT / "business-plan-writer/scripts" / script, ROOT / "pitch-deck-writer/scripts" / script, shallow=False),
    )
style_guides = sorted(ROOT.glob("*/references/writing-style.md"))
check("every skill ships the writing-style guide", len(style_guides) == 5, str(style_guides))
check("writing-style.md copies are identical", all(filecmp.cmp(style_guides[0], g, shallow=False) for g in style_guides))
check(
    "restyling.md copies are identical",
    filecmp.cmp(ROOT / "business-plan-writer/references/restyling.md", ROOT / "pitch-deck-writer/references/restyling.md", shallow=False),
)
em_dash_docs = [str(p.relative_to(ROOT)) for p in ROOT.rglob("*.md") if p.name != "writing-style.md" and "—" in p.read_text()]
check("no em dashes in the skills' own docs", not em_dash_docs, str(em_dash_docs))

with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    r = run(DOCX, FIX / "coffee-plan-spec.json", tmp / "plan.docx", "--facts", FIX / "coffee-facts.json")
    check("business plan builds cleanly", r.returncode == 0 and "WARNING" not in r.stderr, r.stderr)

    r = run(PPTX, FIX / "coffee-deck-spec.json", tmp / "deck.pptx", "--facts", FIX / "coffee-facts.json")
    check("pitch deck builds cleanly", r.returncode == 0 and "WARNING" not in r.stderr, r.stderr)

    from pptx import Presentation  # noqa: E402
    facts = json.loads((FIX / "coffee-facts.json").read_text())
    text = " ".join(sh.text_frame.text for s in Presentation(str(tmp / "deck.pptx")).slides for sh in s.shapes if sh.has_text_frame)
    check("engine IRR appears in the deck", f"{facts['results']['irrPct']:.1f}%" in text)

    bad = json.loads((FIX / "coffee-plan-spec.json").read_text())
    bad["sections"][0]["blocks"].append({"type": "paragraph", "text": "IRR is {{results.irrr|pct}}"})
    (tmp / "bad.json").write_text(json.dumps(bad))
    r = run(DOCX, tmp / "bad.json", tmp / "bad.docx", "--facts", FIX / "coffee-facts.json")
    check("a typo'd engine path stops the build", r.returncode == 2 and "irrr" in r.stderr, r.stderr)

    r = run(DOCX, FIX / "coffee-plan-spec.json", tmp / "nofacts.docx")
    check("engine tokens without a facts file stop the build", r.returncode == 2, r.stderr)

    styled = json.loads((FIX / "coffee-plan-spec.json").read_text())
    styled["sections"][0]["blocks"].append({"type": "paragraph", "text": "This is a pivotal, seamless opportunity — truly."})
    (tmp / "styled.json").write_text(json.dumps(styled))
    r = run(DOCX, tmp / "styled.json", tmp / "styled.docx", "--facts", FIX / "coffee-facts.json")
    check("style check flags em dashes and hype words", r.returncode == 1 and "em dash" in r.stderr and "pivotal" in r.stderr, r.stderr)

    CHECK = ROOT / "business-plan-writer" / "scripts" / "check_numbers.py"
    r = run(CHECK, tmp / "plan.docx", tmp / "plan.docx")
    check("check_numbers passes a file against itself", r.returncode == 0, r.stdout)
    r = run(CHECK, tmp / "deck.pptx", tmp / "deck.pptx")
    check("check_numbers passes a deck against itself", r.returncode == 0, r.stdout)

    from docx import Document  # noqa: E402
    d = Document(str(tmp / "plan.docx"))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if "£97,350" in cell.text:
                    cell.paragraphs[0].runs[0].text = cell.paragraphs[0].runs[0].text.replace("£97,350", "£97,530")
    d.save(str(tmp / "tampered.docx"))
    r = run(CHECK, tmp / "plan.docx", tmp / "tampered.docx")
    check("check_numbers catches a changed figure", r.returncode == 1 and "£97,350" in r.stdout and "£97,530" in r.stdout, r.stdout)

    tpl = Document()
    tpl.styles["Normal"].font.name = "Garamond"
    tpl.sections[0].header.paragraphs[0].text = "ACME Ventures, confidential"
    del_style = tpl.styles["List Bullet"]
    del_style.element.getparent().remove(del_style.element)  # a template missing a style we use
    tpl.add_paragraph("old template body text")
    tpl.save(str(tmp / "tpl.docx"))
    r = run(DOCX, FIX / "coffee-plan-spec.json", tmp / "templated.docx", "--facts", FIX / "coffee-facts.json", "--template", tmp / "tpl.docx")
    out = Document(str(tmp / "templated.docx"))
    body = "\n".join(p.text for p in out.paragraphs)
    check(
        "--template keeps the founder's styles and header and replaces the body",
        r.returncode == 0
        and out.styles["Normal"].font.name == "Garamond"
        and out.sections[0].header.paragraphs[0].text == "ACME Ventures, confidential"
        and "old template body text" not in body
        and "Executive summary" in body,
        r.stderr,
    )

sys.exit(1 if failures else 0)
